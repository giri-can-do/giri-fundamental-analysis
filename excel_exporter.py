from pathlib import Path

from openpyxl import Workbook
from openpyxl import chart
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter


OUTPUT_DIR = Path("exports")


def create_workbook(
    ticker,
    company_name,
    income_statement,
    profitability,
    balance_sheet,
    financial_health,
    cash_flow_analysis,
):

    OUTPUT_DIR.mkdir(exist_ok=True)

    wb = Workbook()

    # Remove default worksheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    dashboard_ws = wb.create_sheet("Dashboard")
    income_ws = wb.create_sheet("Income Statement")
    balance_ws = wb.create_sheet("Balance Sheet")
    cashflow_ws = wb.create_sheet("Cash Flow")

    # build Income Statement sheet
    build_income_statement_sheet(
        income_ws,
        ticker,
        company_name,
        income_statement,
    )

    # build Balance Sheet sheet
    build_balance_sheet(
        balance_ws,
        ticker,
        company_name,
        balance_sheet,
    )

    # build Cash Flow sheet
    build_cash_flow_sheet(
        cashflow_ws,
        ticker,
        company_name,
        cash_flow_analysis,
    )

    # build Dashboard sheet
    build_dashboard(
        dashboard_ws,
        ticker,
        company_name,
        income_statement,
    )

    years = [
        record["end"]
        for record
        in income_statement[
            "revenue"
        ]["records"]
    ]

    income_ws = dashboard_ws.parent[
        "Income Statement"
    ]

    for column, year in enumerate(
        years,
        start=2,
    ):
        fiscal_year = year[:4]

        dashboard_ws.cell(
            row=40,
            column=column,
            value=f"FY{fiscal_year}",
        )

    dashboard_ws.row_dimensions[40].hidden = True

    add_growth_chart(
        dashboard_ws,
        income_ws,
        years,
    )

    add_margin_chart(
        dashboard_ws,
        income_ws,
        years,
    )

    output_file = (
        OUTPUT_DIR
        / f"Fundamental_Analysis_{ticker}.xlsx"
    )

    wb.save(output_file)

    return output_file

def build_income_statement_sheet(
    ws,
    ticker,
    company_name,
    statement,
):

    ws.sheet_view.showGridLines = False

    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    ws["A1"] = "FUNDAMENTAL ANALYSIS"
    ws["A2"] = company_name
    ws["A3"] = ticker

    ws["A1"].font = Font(
        size=18,
        bold=True,
    )

    ws["A2"].font = Font(
        size=14,
        bold=True,
    )

    # --------------------------------------------------
    # YEARS
    # --------------------------------------------------

    revenue_records = statement["revenue"]["records"]

    years = [
        record["end"]
        for record in revenue_records
    ]

    ws["A5"] = "Metric"

    # --------------------------------------------------
    # RAW INCOME STATEMENT
    # --------------------------------------------------

    metric_rows = {
        "revenue": 6,
        "gross_profit": 8,
        "operating_income": 10,
        "net_income": 12,
        "eps": 14,
    }

    growth_rows = {
        "revenue": 7,
        "gross_profit": 9,
        "operating_income": 11,
        "net_income": 13,
        "eps": 15,
    }

    for metric_key, row in metric_rows.items():

        metric = statement[metric_key]

        ws.cell(
            row=row,
            column=1,
            value=metric["label"],
        )

        records_by_date = {
            record["end"]: record["val"]
            for record in metric["records"]
        }

        for column, year in enumerate(
            years,
            start=2,
        ):

            ws.cell(
                row=row,
                column=column,
                value=records_by_date.get(year),
            )

    # --------------------------------------------------
    # YOY GROWTH — EXCEL FORMULAS
    # --------------------------------------------------

    for metric_key, growth_row in growth_rows.items():

        value_row = metric_rows[metric_key]

        ws.cell(
            row=growth_row,
            column=1,
            value="YoY Growth",
        )

        # First year has no previous year
        ws.cell(
            row=growth_row,
            column=2,
            value="N/A",
        )

        for column in range(
            3,
            len(years) + 2,
        ):

            current_cell = (
                f"{get_column_letter(column)}"
                f"{value_row}"
            )

            previous_cell = (
                f"{get_column_letter(column - 1)}"
                f"{value_row}"
            )

            formula = (
                f'=IF('
                f'{previous_cell}=0,'
                f'"N/A",'
                f'({current_cell}-{previous_cell})'
                f'/ABS({previous_cell})'
                f')'
            )

            ws.cell(
                row=growth_row,
                column=column,
                value=formula,
            )

    # --------------------------------------------------
    # PROFITABILITY SECTION
    # --------------------------------------------------

    ws["A18"] = "PROFITABILITY"

    ws["A19"] = "Gross Margin"
    ws["A20"] = "Operating Margin"
    ws["A21"] = "Net Margin"

    for column in range(
        2,
        len(years) + 2,
    ):

        letter = get_column_letter(column)

        # Gross Profit / Revenue
        ws.cell(
            row=19,
            column=column,
            value=f'=IFERROR({letter}8/{letter}6,"N/A")',
        )

        # Operating Income / Revenue
        ws.cell(
            row=20,
            column=column,
            value=f'=IFERROR({letter}10/{letter}6,"N/A")',
        )

        # Net Income / Revenue
        ws.cell(
            row=21,
            column=column,
            value=f'=IFERROR({letter}12/{letter}6,"N/A")',
        )

    format_income_statement(
        ws,
        years,
    )

def format_income_statement(
    ws,
    years,
):

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    section_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9E2F3",
        )
    )

    last_column = len(years) + 1

    # --------------------------------------------------
    # YEAR HEADER
    # --------------------------------------------------

    for column in range(
        1,
        last_column + 1,
    ):

        cell = ws.cell(
            row=5,
            column=column,
        )

        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------
    # FINANCIAL VALUE ROWS
    # --------------------------------------------------

    value_rows = [
        6,
        8,
        10,
        12,
        14,
    ]

    for row in value_rows:

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True
        )

        for column in range(
            1,
            last_column + 1,
        ):

            ws.cell(
                row=row,
                column=column,
            ).border = thin_border

    # --------------------------------------------------
    # GROWTH ROWS
    # --------------------------------------------------

    growth_rows = [
        7,
        9,
        11,
        13,
        15,
    ]

    for row in growth_rows:

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            italic=True,
        )

        for column in range(
            2,
            last_column + 1,
        ):

            ws.cell(
                row=row,
                column=column,
            ).number_format = "0.00%"

    # --------------------------------------------------
    # USD VALUES
    # --------------------------------------------------

    for row in [
        6,
        8,
        10,
        12,
    ]:

        for column in range(
            2,
            last_column + 1,
        ):

            ws.cell(
                row=row,
                column=column,
            ).number_format = (
                '$#,##0.0,,,"B"'
            )

    # EPS
    for column in range(
        2,
        last_column + 1,
    ):

        ws.cell(
            row=14,
            column=column,
        ).number_format = "$0.00"

    # --------------------------------------------------
    # PROFITABILITY SECTION
    # --------------------------------------------------

    for column in range(
        1,
        last_column + 1,
    ):

        cell = ws.cell(
            row=18,
            column=column,
        )

        cell.fill = section_fill

    ws["A18"].font = Font(
        bold=True,
        size=12,
    )

    for row in [
        19,
        20,
        21,
    ]:

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True
        )

        for column in range(
            2,
            last_column + 1,
        ):

            ws.cell(
                row=row,
                column=column,
            ).number_format = "0.00%"

    # --------------------------------------------------
    # COLUMN WIDTH
    # --------------------------------------------------

    ws.column_dimensions["A"].width = 24

    for column in range(
        2,
        last_column + 1,
    ):

        ws.column_dimensions[
            get_column_letter(column)
        ].width = 16

    # Keep labels and years visible
    ws.freeze_panes = "B6"


def build_balance_sheet(
    ws,
    ticker,
    company_name,
    balance_sheet,
):

    ws.sheet_view.showGridLines = False

    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    ws["A1"] = "FUNDAMENTAL ANALYSIS"
    ws["A2"] = company_name
    ws["A3"] = ticker

    ws["A1"].font = Font(
        size=18,
        bold=True,
    )

    ws["A2"].font = Font(
        size=14,
        bold=True,
    )

    # --------------------------------------------------
    # YEARS
    # --------------------------------------------------

    asset_records = balance_sheet[
        "total_assets"
    ]["records"]

    years = [
        record["end"]
        for record in asset_records
    ]

    ws["A5"] = "Metric"

    for column, year in enumerate(
        years,
        start=2,
    ):
        ws.cell(
            row=5,
            column=column,
            value=year,
        )

    # --------------------------------------------------
    # BALANCE SHEET VALUES
    # --------------------------------------------------

    metric_rows = {
        "cash": 6,
        "current_assets": 7,
        "current_liabilities": 8,
        "total_assets": 10,
        "total_liabilities": 11,
        "equity": 12,
    }

    for metric_key, row in metric_rows.items():

        metric = balance_sheet[metric_key]

        ws.cell(
            row=row,
            column=1,
            value=metric["label"],
        )

        records_by_date = {
            record["end"]: record["val"]
            for record in metric["records"]
        }

        for column, year in enumerate(
            years,
            start=2,
        ):

            ws.cell(
                row=row,
                column=column,
                value=records_by_date.get(year),
            )

    # --------------------------------------------------
    # FINANCIAL HEALTH
    # --------------------------------------------------

    ws["A15"] = "FINANCIAL HEALTH"

    ws["A16"] = "Current Ratio"
    ws["A17"] = "Liabilities / Assets"
    ws["A18"] = "Equity Ratio"

    for column in range(
        2,
        len(years) + 2,
    ):

        letter = get_column_letter(column)

        # Current Assets / Current Liabilities
        ws.cell(
            row=16,
            column=column,
            value=(
                f'=IFERROR('
                f'{letter}7/{letter}8,'
                f'"N/A")'
            ),
        )

        # Total Liabilities / Total Assets
        ws.cell(
            row=17,
            column=column,
            value=(
                f'=IFERROR('
                f'{letter}11/{letter}10,'
                f'"N/A")'
            ),
        )

        # Equity / Total Assets
        ws.cell(
            row=18,
            column=column,
            value=(
                f'=IFERROR('
                f'{letter}12/{letter}10,'
                f'"N/A")'
            ),
        )

    # --------------------------------------------------
    # ACCOUNTING CHECK
    # --------------------------------------------------

    ws["A21"] = "DATA VALIDATION"
    ws["A22"] = (
        "Assets - (Liabilities + Equity)"
    )

    for column in range(
        2,
        len(years) + 2,
    ):

        letter = get_column_letter(column)

        ws.cell(
            row=22,
            column=column,
            value=(
                f'={letter}10-'
                f'({letter}11+{letter}12)'
            ),
        )

    format_balance_sheet(
        ws,
        years,
    )

def format_balance_sheet(
    ws,
    years,
):

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    section_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    validation_fill = PatternFill(
        "solid",
        fgColor="E2F0D9",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9E2F3",
        )
    )

    last_column = len(years) + 1

    # --------------------------------------------------
    # YEAR HEADER
    # --------------------------------------------------

    for column in range(
        1,
        last_column + 1,
    ):

        cell = ws.cell(
            row=5,
            column=column,
        )

        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------
    # RAW BALANCE SHEET
    # --------------------------------------------------

    value_rows = [
        6,
        7,
        8,
        10,
        11,
        12,
    ]

    for row in value_rows:

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True
        )

        for column in range(
            1,
            last_column + 1,
        ):

            cell = ws.cell(
                row=row,
                column=column,
            )

            cell.border = thin_border

            if column >= 2:

                cell.number_format = (
                    '$#,##0.0,,,"B"'
                )

    # --------------------------------------------------
    # FINANCIAL HEALTH SECTION
    # --------------------------------------------------

    for column in range(
        1,
        last_column + 1,
    ):

        ws.cell(
            row=15,
            column=column,
        ).fill = section_fill

    ws["A15"].font = Font(
        bold=True,
        size=12,
    )

    for row in [
        16,
        17,
        18,
    ]:

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True
        )

    # Current Ratio
    for column in range(
        2,
        last_column + 1,
    ):

        ws.cell(
            row=16,
            column=column,
        ).number_format = "0.00"

    # Percentage ratios
    for row in [
        17,
        18,
    ]:

        for column in range(
            2,
            last_column + 1,
        ):

            ws.cell(
                row=row,
                column=column,
            ).number_format = "0.00%"

    # --------------------------------------------------
    # DATA VALIDATION SECTION
    # --------------------------------------------------

    for column in range(
        1,
        last_column + 1,
    ):

        ws.cell(
            row=21,
            column=column,
        ).fill = validation_fill

    ws["A21"].font = Font(
        bold=True,
        size=12,
    )

    ws["A22"].font = Font(
        bold=True
    )

    for column in range(
        2,
        last_column + 1,
    ):

        ws.cell(
            row=22,
            column=column,
        ).number_format = (
            '$#,##0.00,,,"B"'
        )

    # --------------------------------------------------
    # WIDTH / FREEZE
    # --------------------------------------------------

    ws.column_dimensions["A"].width = 34

    for column in range(
        2,
        last_column + 1,
    ):

        ws.column_dimensions[
            get_column_letter(column)
        ].width = 16

    ws.freeze_panes = "B6"

    for column, year in enumerate(
        years,
        start=2,
    ):
        fiscal_year = year[:4]

        ws.cell(
            row=40,
            column=column,
            value=f"FY{fiscal_year}",
        )

    ws.row_dimensions[40].hidden = True

def build_cash_flow_sheet(
    ws,
    ticker,
    company_name,
    cash_flow_analysis,
):

    ws.sheet_view.showGridLines = False

    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    ws["A1"] = "FUNDAMENTAL ANALYSIS"
    ws["A2"] = company_name
    ws["A3"] = ticker

    ws["A1"].font = Font(
        size=18,
        bold=True,
    )

    ws["A2"].font = Font(
        size=14,
        bold=True,
    )

    # --------------------------------------------------
    # YEARS
    # --------------------------------------------------

    years = [
        row["period_end"]
        for row in cash_flow_analysis
    ]

    ws["A5"] = "Metric"

    for column, year in enumerate(
        years,
        start=2,
    ):
        ws.cell(
            row=5,
            column=column,
            value=year,
        )

    # --------------------------------------------------
    # RAW CASH FLOW
    # --------------------------------------------------

    ws["A6"] = "Operating Cash Flow"
    ws["A7"] = "Capital Expenditures"

    for column, row in enumerate(
        cash_flow_analysis,
        start=2,
    ):

        ws.cell(
            row=6,
            column=column,
            value=row["operating_cash_flow"],
        )

        ws.cell(
            row=7,
            column=column,
            value=row["capex"],
        )

    # --------------------------------------------------
    # CASH FLOW ANALYSIS
    # --------------------------------------------------

    ws["A10"] = "CASH FLOW ANALYSIS"

    ws["A11"] = "Free Cash Flow"
    ws["A12"] = "FCF Margin"
    ws["A13"] = "Cash Conversion"

    for column in range(
        2,
        len(years) + 2,
    ):

        letter = get_column_letter(column)

        # Free Cash Flow
        # Operating CF - CapEx
        ws.cell(
            row=11,
            column=column,
            value=f"={letter}6-{letter}7",
        )

        # FCF Margin
        # FCF / Revenue
        ws.cell(
            row=12,
            column=column,
            value=(
                f'=IFERROR('
                f'{letter}11/'
                f'\'Income Statement\'!{letter}6,'
                f'"N/A")'
            ),
        )

        # Cash Conversion
        # Operating CF / Net Income
        ws.cell(
            row=13,
            column=column,
            value=(
                f'=IFERROR('
                f'{letter}6/'
                f'\'Income Statement\'!{letter}12,'
                f'"N/A")'
            ),
        )

    format_cash_flow_sheet(
        ws,
        years,
    )

def format_cash_flow_sheet(
    ws,
    years,
):

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    section_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9E2F3",
        )
    )

    last_column = len(years) + 1

    # --------------------------------------------------
    # HEADER
    # --------------------------------------------------

    for column in range(
        1,
        last_column + 1,
    ):

        cell = ws.cell(
            row=5,
            column=column,
        )

        cell.fill = header_fill
        cell.font = white_font

        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------
    # RAW VALUES
    # --------------------------------------------------

    for row in [
        6,
        7,
    ]:

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True
        )

        for column in range(
            1,
            last_column + 1,
        ):

            cell = ws.cell(
                row=row,
                column=column,
            )

            cell.border = thin_border

            if column >= 2:

                cell.number_format = (
                    '$#,##0.0,,,"B"'
                )

    # --------------------------------------------------
    # ANALYSIS SECTION
    # --------------------------------------------------

    for column in range(
        1,
        last_column + 1,
    ):

        ws.cell(
            row=10,
            column=column,
        ).fill = section_fill

    ws["A10"].font = Font(
        bold=True,
        size=12,
    )

    for row in [
        11,
        12,
        13,
    ]:

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True
        )

    # FCF
    for column in range(
        2,
        last_column + 1,
    ):

        ws.cell(
            row=11,
            column=column,
        ).number_format = (
            '$#,##0.0,,,"B"'
        )

    # FCF Margin + Cash Conversion
    for row in [
        12,
        13,
    ]:

        for column in range(
            2,
            last_column + 1,
        ):

            ws.cell(
                row=row,
                column=column,
            ).number_format = "0.00%"

    # --------------------------------------------------
    # WIDTH / FREEZE
    # --------------------------------------------------

    ws.column_dimensions["A"].width = 28

    for column in range(
        2,
        last_column + 1,
    ):

        ws.column_dimensions[
            get_column_letter(column)
        ].width = 16

    ws.freeze_panes = "B6"

def build_dashboard(
    ws,
    ticker,
    company_name,
    income_statement,
):

    ws.sheet_view.showGridLines = False

    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    ws.merge_cells("A1:L1")
    ws["A1"] = (
        f"FUNDAMENTAL ANALYSIS — "
        f"{company_name} ({ticker})"
    )

    ws["A1"].font = Font(
        size=20,
        bold=True,
        color="FFFFFF",
    )

    ws["A1"].fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    ws["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[1].height = 32

    latest_year = (
        income_statement["revenue"]
        ["records"][-1]["end"]
    )

    latest_column_number = (
        len(
            income_statement[
                "revenue"
            ]["records"]
        )
        + 1
    )

    latest_column = get_column_letter(
        latest_column_number
    )

    ws.merge_cells("A2:L2")

    ws["A2"] = (
        f"Latest Fiscal Year: {latest_year}"
    )

    ws["A2"].font = Font(
        italic=True,
        color="666666",
    )

    # --------------------------------------------------
    # GROWTH
    # --------------------------------------------------

    create_dashboard_section(
        ws,
        row=4,
        title="GROWTH",
    )

    create_kpi(
        ws,
        label_cell="A5",
        value_cell="A6",
        label="Revenue Growth",
        formula=(
            f"='Income Statement'!"
            f"{latest_column}7"
        ),
        number_format="0.00%",
    )

    create_kpi(
        ws,
        label_cell="E5",
        value_cell="E6",
        label="EPS Growth",
        formula=(
            f"='Income Statement'!"
            f"{latest_column}15"
        ),
        number_format="0.00%",
    )

    create_kpi(
        ws,
        label_cell="I5",
        value_cell="I6",
        label="Net Income Growth",
        formula=(
            f"='Income Statement'!"
            f"{latest_column}13"
        ),
        number_format="0.00%",
    )

    # --------------------------------------------------
    # PROFITABILITY
    # --------------------------------------------------

    create_dashboard_section(
        ws,
        row=8,
        title="PROFITABILITY",
    )

    create_kpi(
        ws,
        "A9",
        "A10",
        "Gross Margin",
        (
            f"='Income Statement'!"
            f"{latest_column}19"
        ),
        "0.00%",
    )

    create_kpi(
        ws,
        "E9",
        "E10",
        "Operating Margin",
        (
            f"='Income Statement'!"
            f"{latest_column}20"
        ),
        "0.00%",
    )

    create_kpi(
        ws,
        "I9",
        "I10",
        "Net Margin",
        (
            f"='Income Statement'!"
            f"{latest_column}21"
        ),
        "0.00%",
    )

    # --------------------------------------------------
    # FINANCIAL HEALTH
    # --------------------------------------------------

    create_dashboard_section(
        ws,
        row=12,
        title="FINANCIAL HEALTH",
    )

    create_kpi(
        ws,
        "A13",
        "A14",
        "Current Ratio",
        (
            f"='Balance Sheet'!"
            f"{latest_column}16"
        ),
        "0.00",
    )

    create_kpi(
        ws,
        "E13",
        "E14",
        "Liabilities / Assets",
        (
            f"='Balance Sheet'!"
            f"{latest_column}17"
        ),
        "0.00%",
    )

    create_kpi(
        ws,
        "I13",
        "I14",
        "Equity Ratio",
        (
            f"='Balance Sheet'!"
            f"{latest_column}18"
        ),
        "0.00%",
    )

    # --------------------------------------------------
    # CASH GENERATION
    # --------------------------------------------------

    create_dashboard_section(
        ws,
        row=16,
        title="CASH GENERATION",
    )

    create_kpi(
        ws,
        "A17",
        "A18",
        "Operating Cash Flow",
        (
            f"='Cash Flow'!"
            f"{latest_column}6"
        ),
        '$#,##0.0,,,"B"',
    )

    create_kpi(
        ws,
        "E17",
        "E18",
        "Free Cash Flow",
        (
            f"='Cash Flow'!"
            f"{latest_column}11"
        ),
        '$#,##0.0,,,"B"',
    )

    create_kpi(
        ws,
        "I17",
        "I18",
        "FCF Margin",
        (
            f"='Cash Flow'!"
            f"{latest_column}12"
        ),
        "0.00%",
    )

    format_dashboard(ws)


def create_kpi(
    ws,
    label_cell,
    value_cell,
    label,
    formula,
    number_format,
):

    ws[label_cell] = label

    ws[label_cell].font = Font(
        bold=True,
        color="666666",
    )

    ws[value_cell] = formula

    ws[value_cell].font = Font(
        size=18,
        bold=True,
    )

    ws[value_cell].number_format = (
        number_format
    )

def create_dashboard_section(
    ws,
    row,
    title,
):

    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=12,
    )

    cell = ws.cell(
        row=row,
        column=1,
    )

    cell.value = title

    cell.font = Font(
        bold=True,
        color="1F1F1F",
    )

    cell.fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

def format_dashboard(ws):

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 5
    ws.column_dimensions["D"].width = 5

    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 5
    ws.column_dimensions["G"].width = 5
    ws.column_dimensions["H"].width = 5

    ws.column_dimensions["I"].width = 22
    ws.column_dimensions["J"].width = 5
    ws.column_dimensions["K"].width = 5
    ws.column_dimensions["L"].width = 5

    for row in [
        6,
        10,
        14,
        18,
    ]:

        ws.row_dimensions[row].height = 28

    ws.freeze_panes = "A4"

def add_growth_chart(
    dashboard_ws,
    income_ws,
    years,
):

    chart = LineChart()

    chart.title = "Revenue & Net Income Trend"
    chart.style = 10

    chart.height = 6.5
    chart.width = 13

    chart.y_axis.numFmt = '$0,,,"B"'
    chart.x_axis.tickLblPos = "low"
    chart.legend.position = "b"

    # Revenue = row 6
    revenue_data = Reference(
        income_ws,
        min_col=1,
        max_col=len(years) + 1,
        min_row=6,
        max_row=6,
    )

    # Net Income = row 12
    net_income_data = Reference(
        income_ws,
        min_col=1,
        max_col=len(years) + 1,
        min_row=12,
        max_row=12,
    )

    categories = Reference(
        income_ws,
        min_col=2,
        max_col=len(years) + 1,
        min_row=40,
        max_row=40,
    )

    chart.add_data(
        revenue_data,
        titles_from_data=True,
        from_rows=True,
    )

    chart.add_data(
        net_income_data,
        titles_from_data=True,
        from_rows=True,
    )

    chart.set_categories(categories)

    chart.legend.position = "b"

    dashboard_ws.add_chart(
        chart,
        "A21",
    )

def add_margin_chart(
    dashboard_ws,
    income_ws,
    years,
):

    chart = LineChart()

    chart.title = "Profitability Margin Trend"
    chart.style = 10

    chart.height = 6.5
    chart.width = 13

    chart.y_axis.numFmt = "0%"
    chart.x_axis.tickLblPos = "low"

    chart.legend.position = "b"

    categories = Reference(
        income_ws,
        min_col=2,
        max_col=len(years) + 1,
        min_row=40,
        max_row=40,
    )

    # Gross Margin
    gross_margin = Reference(
        income_ws,
        min_col=1,
        max_col=len(years) + 1,
        min_row=19,
        max_row=19,
    )

    # Operating Margin
    operating_margin = Reference(
        income_ws,
        min_col=1,
        max_col=len(years) + 1,
        min_row=20,
        max_row=20,
    )

    # Net Margin
    net_margin = Reference(
        income_ws,
        min_col=1,
        max_col=len(years) + 1,
        min_row=21,
        max_row=21,
    )

    chart.add_data(
        gross_margin,
        titles_from_data=True,
        from_rows=True,
    )

    chart.add_data(
        operating_margin,
        titles_from_data=True,
        from_rows=True,
    )

    chart.add_data(
        net_margin,
        titles_from_data=True,
        from_rows=True,
    )

    chart.set_categories(categories)

    chart.legend.position = "b"

    # Excel percentage scale
    chart.y_axis.numFmt = "0%"

    dashboard_ws.add_chart(
        chart,
        "G21",
    )